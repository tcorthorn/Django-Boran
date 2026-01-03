from django.contrib import admin
from django.urls import path
from django.shortcuts import redirect
from django.contrib import messages
from boran_app.models import VentasConsulta
from boran_app.utils import regenerar_ventas_consulta
from boran_app.models import MovimientoUnificadoCredito, MovimientoUnificadoDebito, AjusteInventario
from boran_app.models import ResumenCredito, ResumenDebito
from .models import Inventario, InvEP, InvVP, InvEPVP
from boran_app.utils import regenerar_resumenes_credito_debito
# Register your models here.
from .models import AsientosContables, Envios, EntradaProductos, Catalogo
from .models import OtrosGastos, SueldosHonorarios, BalanceInicial, InventarioInicial, Ventas, VentasConsulta
#from .models import InventarioActual
from django.db import models
from boran_app.admin_export_excel_mixin import ExportExcelMixin
from django import forms
from django.core.exceptions import ValidationError as DjangoValidationError

# --- Helper para normalizar cualquier ValidationError (evitar KeyError 'model') ---
def _normalize_validation_error(e):
    """
    Recibe un DjangoValidationError y devuelve lista de strings sin placeholders rotos.
    """
    msgs = []
    # e puede ser dict (por campo) o lista (no-field)
    if hasattr(e, 'error_dict') and e.error_dict:
        for field, errs in e.error_dict.items():
            for err in errs:
                msg = getattr(err, 'message', str(err))
                params = getattr(err, 'params', {}) or {}
                # Intentar interpolar si hay params; si faltan, dejar tal cual para no romper
                try:
                    msg = msg % params
                except Exception:
                    pass
                if field != '__all__':
                    msg = f"{field}: {msg}"
                msgs.append(msg)
    else:
        for err in getattr(e, 'error_list', [e]):
            msg = getattr(err, 'message', str(err))
            params = getattr(err, 'params', {}) or {}
            try:
                msg = msg % params
            except Exception:
                pass
            msgs.append(msg)
    return msgs


@admin.register(AsientosContables)
class AsientosContablesAdmin(ExportExcelMixin, admin.ModelAdmin):
    list_display = ('id', 'fecha', 'monto', 'debito', 'cuenta_debito', 'credito', 'cuenta_credito','comentario')
    #list_filter = ('fecha','cuenta_debito','cuenta_credito')
    search_fields = ('cuenta_debito', 'cuenta_credito')
    ordering = ('-fecha',)  # más reciente → más antigua
    # Opcionalmente, puedes definir campos a exportar:
    # export_excel_fields = ('id', 'fecha', 'monto', ...)

@admin.register(Envios)
class EnviosAdmin(ExportExcelMixin, admin.ModelAdmin):
    list_display = ('id', 'fecha', 'tienda_bodega', 'sku', 'cantidad')
    list_filter = ('tienda_bodega', 'fecha')
    search_fields = ('sku',)
    ordering = ('-fecha',)  # más reciente → más antigua

# boran_app/admin.py (o consult_app/admin.py, donde tengas tu admin)

class EntradaProductosAdminForm(forms.ModelForm):
    class Meta:
        model = EntradaProductos
        fields = '__all__'

    def clean(self):
        try:
            return super().clean()
        except DjangoValidationError as e:
            # Convierte cualquier ValidationError con placeholders a mensajes de texto plano
            msgs = []
            if hasattr(e, 'error_dict') and e.error_dict:
                for field, errs in e.error_dict.items():
                    for err in errs:
                        msg = getattr(err, 'message', str(err))
                        params = getattr(err, 'params', {}) or {}
                        try:
                            msg = msg % params
                        except Exception:
                            pass
                        if field != '__all__':
                            msg = f"{field}: {msg}"
                        msgs.append(msg)
            else:
                for err in getattr(e, 'error_list', [e]):
                    msg = getattr(err, 'message', str(err))
                    params = getattr(err, 'params', {}) or {}
                    try:
                        msg = msg % params
                    except Exception:
                        pass
                    msgs.append(msg)
            raise forms.ValidationError(msgs)




@admin.register(EntradaProductos)
class EntradaProductosAdmin(ExportExcelMixin,admin.ModelAdmin):
    form = EntradaProductosAdminForm   # 👈 agrega esto
    # Mostrar columnas en la lista de registros
    list_display = (
        'fecha',
        'sku',
        'cantidad_ingresada',
        'costo_con_iva',
        'costo_adicional',
        'costo_unitario_total_producto',
        'costo_total',
        'iva_compra',
        'costo_neto',
        'debito',
        'debito_iva',
        'cuenta_credito',
        'credito',
        'inventario_inicio',
    )
    list_filter = ('sku__sku', 'fecha')
    search_fields = ('sku__sku',)  # permite buscar por valor de SKU
    ordering = ('-fecha',)  # más reciente → más antigua

    # Para que los campos calculados aparezcan en el formulario como sólo lectura:
    readonly_fields = (
        'costo_unitario_total_producto',
        'costo_total',
        'iva_compra',
        'costo_neto',
        'debito',
        'debito_iva',
        'cuenta_credito',
        'credito',
    )

    # Controlar el orden de todos los campos en el formulario de edición/creación
    fields = (
        'fecha',
        'sku',
        'cantidad_ingresada',
        'costo_con_iva',
        'costo_adicional',
        'costo_unitario_total_producto',
        'costo_total',
        'iva_compra',
        'costo_neto',
        'cuenta_debito',
        'debito',
        'cuenta_debito_iva',
        'debito_iva',
        'inventario_inicio',
        'cuenta_credito',
        'credito',
        'comentario',
    )

    def _changeform_view(self, request, object_id=None, form_url='', extra_context=None):
        """
        Sobrescribe el flujo default para capturar VALIDATION ERRORS tardíos
        (p. ej. lanzados en save(), señales o constraints) y normalizarlos,
        evitando el KeyError 'model' al renderizar.
        """
        try:
            return super()._changeform_view(request, object_id, form_url, extra_context)
        except DjangoValidationError as e:
            # Reconstruye el form con los datos del POST y añade errores “seguros”
            ModelForm = self.get_form(request)
            form = ModelForm(request.POST, request.FILES)
            for msg in _normalize_validation_error(e):
                form.add_error(None, msg)  # errores no-campo
            context = {
                **(extra_context or {}),
                'add': object_id is None,
                'change': object_id is not None,
                'form_url': form_url,
                'form': form,
                'opts': self.model._meta,
                'original': None,
                'title': "Agregar Entrada de Producto" if object_id is None else "Modificar Entrada de Producto",
            }
            return self.render_change_form(request, context, add=object_id is None, change=object_id is not None, obj=None)

@admin.register(Catalogo)
class CatalogoAdmin(ExportExcelMixin,admin.ModelAdmin):
    # Aquí defines qué campos mostrar en la lista o en el formulario de edición
    list_display = ['sku', 'producto', 'categoria', 'fecha_ingreso','costo_base', 'costo_adicional', 'costo_directo_producto']
    #list_editable = ['producto','categoria'] # Aqui para editar directo de la lista del Admin
    search_fields = ('sku', 'producto', 'categoria')
    list_filter = ('categoria','producto','sku')
   
    


@admin.register(OtrosGastos)
class OtrosGastosAdmin(ExportExcelMixin, admin.ModelAdmin):
    # Estas opciones son sugeridas; ajústalas según lo que quieras ver en la lista
    list_display = ('id','fecha', 'otros_gastos', 'total', 'iva','cuenta_debito', 'monto_neto',  'cuenta_credito','credito','debito_eerr','cuenta_debito_eerr','comentario')
    search_fields = ('otros_gastos', 'comentario')
    list_filter = ('otros_gastos', 'fecha','cuenta_debito','cuenta_credito','cuenta_debito_eerr')
    ordering = ('-fecha',)  # más reciente → más antigua


@admin.register(SueldosHonorarios)
class SueldosHonorariosAdmin(ExportExcelMixin,admin.ModelAdmin):
    list_display = ('fecha', 'tipo_remuneracion', 'nombre', 'monto_total_pagado')
    list_filter = ('tipo_remuneracion', 'fecha','nombre')
    ordering = ('-fecha',)  # más reciente → más antigua
    readonly_fields = (
        'debito', 'credito', 'credito2', 
        'cuenta_debito', 'cuenta_credito', 'cuenta_credito2',
        'retenciones'
    )

@admin.register(BalanceInicial)
class BalanceInicialAdmin(ExportExcelMixin, admin.ModelAdmin):
    list_display = ('fecha', 'cuenta_debito', 'debito', 'cuenta_credito', 'credito', 'comentario')
    search_fields = ('comentario',)
    list_filter = ('fecha',)
    ordering = ('-fecha',)  # más reciente → más antigua

@admin.register(InventarioInicial)
class InventarioInicialAdmin(ExportExcelMixin, admin.ModelAdmin):
    list_display = ('sku', 'categoria', 'producto', 'stock', 'bodega', 'comentario')
    search_fields = ('sku', 'producto', 'categoria')
    list_filter = ('categoria', 'bodega')

@admin.register(Ventas)
class VentasAdmin(ExportExcelMixin, admin.ModelAdmin):
    list_display = (
        'id',
        'fecha',
        'numero_pedido',
        'comprador',
        'sku',
        'sku__producto',
        'sku__categoria',
        'cantidad',
        'valor_unitario_venta',
        'valor_envio_cobrado',
        'costo_unitario_venta',
        'total_venta',
        'costo_venta',
        'iva',
        'venta_neta_de_iva',
        'iva_envio',
        'debito',
        'cuenta_credito',
        'credito',
        'debito_eerr',
        'cuenta_credito_eerr',
        'credito_eerr',
        'costo_directo',
        'credito_iva',
        'cuenta_credito_iva',
        'debito_envio',
        'credito_envio',
        'venta_bruta',
        'comision_flow',
        'comision_plataformas_pago',
    )
    readonly_fields = (
        'total_venta',
        'iva_calculo',
        'iva',
        'venta_neta_de_iva',
        'iva_envio',
        'cuenta_debito',
        'debito',
        'debito_eerr',
        'credito_eerr',
        'credito_iva',
        'cuenta_credito_iva',
        'debito_envio',
        'credito_envio',
        'venta_bruta',
        'comision_flow',
        'comision_plataformas_pago',
    )
    list_filter = ('fecha', 'comprador', 'sku__sku')
    search_fields = ('numero_pedido', 'comprador', 'sku__sku', 'sku__producto', 'sku__categoria')
    ordering = ('-fecha',)  # más reciente → más antigua

@admin.register(VentasConsulta)
class VentasConsultaAdmin(ExportExcelMixin,admin.ModelAdmin):
    list_display = ['fecha', 'codigo_producto', 'categoria', 'producto','comprador', 'cantidad','total_venta', 'venta_neta_iva','costo_venta','costo_directo_producto']
    search_fields = ('comprador', 'codigo_producto','categoria', 'producto','comprador', 'cantidad','total_venta', 'costo_venta')
    actions = ['regenerar_tabla']

    def regenerar_tabla(self, request, queryset):
        total = regenerar_ventas_consulta()
        self.message_user(request, f"✅ Regenerado con éxito: {total} registros.", level=messages.SUCCESS)

    regenerar_tabla.short_description = "🔁 Regenerar tabla VentasConsulta"

@admin.register(MovimientoUnificadoCredito)
class MovimientoUnificadoCreditoAdmin(ExportExcelMixin,admin.ModelAdmin):
    list_display = ['fecha', 'cta_credito', 'monto_credito', 'tabla_origen']
    list_filter = ['tabla_origen','cta_credito']
    search_fields = ['texto_coment']

@admin.register(MovimientoUnificadoDebito)
class MovimientoUnificadoDebitoAdmin(ExportExcelMixin,admin.ModelAdmin):
    list_display = ['fecha', 'cta_debito', 'monto_debito', 'tabla_origen']
    list_filter = ['tabla_origen']
    search_fields = ['texto_coment']

@admin.register(ResumenCredito)
class ResumenCreditoAdmin(ExportExcelMixin,admin.ModelAdmin):
    list_display = ['cuenta_credito', 'total_credito']
    ordering = ['cuenta_credito']
    actions = ['regenerar_resumen']

    def regenerar_resumen(self, request, queryset):
        total_creditos, total_debitos = regenerar_resumenes_credito_debito()
        self.message_user(
            request,
            f"✅ Resúmenes regenerados: {total_creditos} créditos y {total_debitos} débitos.",
            level=messages.SUCCESS
        )
    regenerar_resumen.short_description = "🔁 Regenerar Resumen Crédito/Débito"

@admin.register(ResumenDebito)
class ResumenDebitoAdmin(ExportExcelMixin,admin.ModelAdmin):
    list_display = ['cuenta_debito', 'total_debito']
    ordering = ['cuenta_debito']

@admin.register(AjusteInventario)
class AjusteInventarioAdmin(ExportExcelMixin,admin.ModelAdmin):
    list_display = ('fecha', 'sku', 'cantidad', 'costo_producto', 'cuenta_debito', 'debito', 'cuenta_credito', 'comentario')
    search_fields = ('sku__sku', 'comentario')

#@admin.register(Inventario)
#class InventarioAdmin(admin.ModelAdmin):
    #list_display = ('cod_producto', 'producto', 'categoria', 'ingresado', 'vendido', 'ajuste')
    #search_fields = ('cod_producto', 'producto', 'categoria')

#admin.site.register(InvEP)
#admin.site.register(InvVP)
#admin.site.register(InvEPVP)

# ——————————————————————————————————————————————————————————————
#  INVENTARIO ACTUAL
# ——————————————————————————————————————————————————————————————

from django.contrib import admin
from django.db.models import Sum, F, Value, IntegerField
from django.http import HttpResponse
from openpyxl import Workbook


from .models import (
    Catalogo, InventarioInicial, EntradaProductos,
    Envios, Ventas, AjusteInventario, inventarioactualproxy
)

from django.urls import path  # Añade esto si no lo tienes
from django.db import models  # Asegúrate de importar models

@admin.register(inventarioactualproxy)
class inventarioactualadmin(admin.ModelAdmin):
    list_display = (
        'sku', 'categoria', 'producto',
        'stock_display', 'bodega_display',
        'ingresos_display', 'envios_display',
        'ventas_display', 'ajustes_display',
        'en_oficina_display', 'en_bodega_display', 'total_display',  # <-- Nuevo
    )
    search_fields = ('sku', 'categoria', 'producto')
    change_list_template = "admin/boran_app/inventarioactualproxy/change_list.html"

    # Tus métodos originales...
    def stock_display(self, obj):
        ini = InventarioInicial.objects.filter(sku=obj.sku).first()
        return ini.stock if ini else 0
    stock_display.short_description = "Inicial"

    def bodega_display(self, obj):
        ini = InventarioInicial.objects.filter(sku=obj.sku).first()
        return ini.bodega if ini else 0
    bodega_display.short_description = "Bodega"

    def ingresos_display(self, obj):
        return EntradaProductos.objects.filter(sku__sku=obj.sku).aggregate(total=models.Sum('cantidad_ingresada'))['total'] or 0
    ingresos_display.short_description = "Ingresos"

    def envios_display(self, obj):
        return Envios.objects.filter(sku__sku=obj.sku).aggregate(total=models.Sum('cantidad'))['total'] or 0
    envios_display.short_description = "Envios"

    def ventas_display(self, obj):
        return Ventas.objects.filter(sku__sku=obj.sku).aggregate(total=models.Sum('cantidad'))['total'] or 0
    ventas_display.short_description = "Ventas"

    def ajustes_display(self, obj):
        return AjusteInventario.objects.filter(sku__sku=obj.sku).aggregate(total=models.Sum('cantidad'))['total'] or 0
    ajustes_display.short_description = "Ajustes"

    # ——— NUEVOS MÉTODOS CALCULADOS ———
    def en_oficina_display(self, obj):
        inicial = self.stock_display(obj)
        ingresos = self.ingresos_display(obj)
        envios = self.envios_display(obj)
        return inicial + ingresos - envios
    en_oficina_display.short_description = "En Oficina"

    def en_bodega_display(self, obj):
        bodega = self.bodega_display(obj)
        envios = self.envios_display(obj)
        ventas = self.ventas_display(obj)
        return bodega + envios - ventas
    en_bodega_display.short_description = "En Bodega"

    def total_display(self, obj):
        en_oficina = self.en_oficina_display(obj)
        en_bodega = self.en_bodega_display(obj)
        ajustes = self.ajustes_display(obj)
        return en_oficina + en_bodega - ajustes
    total_display.short_description = "Total"

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'exportar_inventario_actual/',
                self.admin_site.admin_view(self.exportar_excel_view),
                name='exportar_inventario_actual'
            ),
        ]
        return custom_urls + urls

    def exportar_excel_view(self, request):
        queryset = Catalogo.objects.all()
        wb = Workbook()
        ws = wb.active
        ws.title = "Inventario Actual"
        ws.append([
            'SKU', 'Categoría', 'Producto', 'Stock', 'Bodega',
            'Ingresos', 'Envios', 'Ventas', 'Ajustes',
            'En Oficina', 'En Bodega', 'Total'
        ])
        for obj in queryset:
            ws.append([
                obj.sku, obj.categoria, obj.producto,
                self.stock_display(obj), self.bodega_display(obj),
                self.ingresos_display(obj), self.envios_display(obj),
                self.ventas_display(obj), self.ajustes_display(obj),
                self.en_oficina_display(obj), self.en_bodega_display(obj),
                self.total_display(obj),
            ])
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=InventarioActual.xlsx'
        wb.save(response)
        return response